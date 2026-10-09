"""tape_sheet.py - "Matchup" sheet in NCAA_FBS_Teams.xlsm: the Excel twin of Tale of the Tape (Lucas 10/9).

Three dropdowns (Team A, Team B, site) and nothing but formulas underneath - no macros, so it works
the moment the workbook opens. Hidden data sheets carry the same numbers as tape_data_2026.json
(the artifact's file), written fresh on every refresh:
    _Tape        one row per rated team: rating / July / rank / QB-out, efficiency shadow split,
                 resume summary (luck, best, worst), season stats with FBS ranks
    _TapeGames   one row per team-game (key "Team|n"), graded against today's ratings
    _TapeCurve   margin -> home win probability (the production curve at the on-air sd)
    _TapeCommon  common opponents per pair (key "A|B"), only pairs that have any

Layout (Matchup): verdict block, 1 The number, 2 The plays (SHADOW), 3 The tape, 4 The resume.
NO MARKET NUMBERS. The score line is a pace guess (season scoring averages), not the on-air call.

    python tape_sheet.py            # rebuild tape_data_2026.json, write the sheets, recalc via Excel
    python tape_sheet.py --no-data  # reuse the json on disk
refresh_all.py calls write_tape_sheet(book) after the team-stats sheet.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import DataBarRule, CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
DATA = HERE / "tape_data_2026.json"
BOOK = HERE / "NCAA_FBS_Teams.xlsm"
SLOTS = 13          # resume rows per team (12 regular-season games + a spare)
SITES = ["B at home", "Neutral", "A at home"]

ARIAL = Font(name="Arial", size=10)
ARIAL_B = Font(name="Arial", size=10, bold=True)
MUTED = Font(name="Arial", size=9, italic=True, color="5C6B7E")
TITLE = Font(name="Arial", size=16, bold=True)
BIG = Font(name="Arial", size=22, bold=True)
WHITE_B = Font(name="Arial", size=11, bold=True, color="FFFFFF")
PICK_FILL = PatternFill("solid", start_color="FFF2CC")
HEAD_FILL = PatternFill("solid", start_color="1F3864")
SOFT_FILL = PatternFill("solid", start_color="F2F2F2")
SHADOW_FILL = PatternFill("solid", start_color="5B4DB1")
THIN = Side(style="thin", color="D9D9D9")


def q(s: str) -> str:
    """Excel string literal."""
    return '"' + s.replace('"', '""') + '"'


# ---------------- hidden data sheets ----------------

def _hidden(wb, name):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_state = "hidden"
    return ws


def write_data_sheets(wb, d: dict) -> dict:
    """Returns the column map for _Tape (field -> column letter)."""
    keys = [k for k, _, _ in d["stat_keys"]]
    fields = ["key", "name", "abbr", "conf", "record", "rated_gp", "july", "now", "delta", "rank", "qb_out",
              "eff_off", "eff_def", "eff_st", "eff_total", "eff_rank",
              "luck", "best_opp", "best_resid", "worst_opp", "worst_resid", "resid_sum", "rated_n", "games_n"]
    fields += [f"off_{k}" for k in keys] + [f"off_{k}_rk" for k in keys] + [f"def_{k}" for k in keys] + [f"def_{k}_rk" for k in keys]
    col = {f: L(i + 1) for i, f in enumerate(fields)}

    ws = _hidden(wb, "_Tape")
    for j, f in enumerate(fields, 1):
        ws.cell(row=1, column=j, value=f)
    teams = sorted(d["teams"].values(), key=lambda t: t["name"])
    to_pts = d["rule"]["to_pts"]
    for i, t in enumerate(teams, 2):
        rated = [g for g in t["games"] if g["rated"]]
        best = max(rated, key=lambda g: g["resid"]) if rated else None
        worst = min(rated, key=lambda g: g["resid"]) if rated else None
        e = t["eff"] or {}
        st = t["stats"] or {"off": {}, "def_": {}}
        row = dict(key=t["key"], name=t["name"], abbr=t["abbr"], conf=t["conf"], record=t["record"], rated_gp=t["rated_gp"],
                   july=t["july"], now=t["now"], delta=t["delta"], rank=t["rank"], qb_out=t["qb_out"] or 0,
                   eff_off=e.get("off"), eff_def=e.get("def_"), eff_st=e.get("st"), eff_total=e.get("total"), eff_rank=e.get("rank"),
                   luck=round(sum(g["to"] for g in rated) * to_pts, 1), best_opp=(best["opp"] if best else ""), best_resid=(best["resid"] if best else None),
                   worst_opp=(worst["opp"] if worst else ""), worst_resid=(worst["resid"] if worst else None),
                   resid_sum=round(sum(g["resid"] for g in rated), 1), rated_n=len(rated), games_n=len(t["games"]))
        for k in keys:
            row[f"off_{k}"] = st["off"].get(k); row[f"off_{k}_rk"] = st["off"].get(k + "_rk")
            row[f"def_{k}"] = st["def_"].get(k); row[f"def_{k}_rk"] = st["def_"].get(k + "_rk")
        for j, f in enumerate(fields, 1):
            ws.cell(row=i, column=j, value=row.get(f))

    wg = _hidden(wb, "_TapeGames")
    for j, h in enumerate(["key", "wk", "opp", "site", "result", "expected", "delucked", "to", "resid", "rated", "opp_rank", "opp_label"], 1):
        wg.cell(row=1, column=j, value=h)
    r = 2
    for t in teams:
        for n, g in enumerate(t["games"], 1):
            site = {"home": "vs", "away": "at", "neutral": "n."}[g["site"]]
            label = f"{site} {g['opp']}" + (f" (#{g['opp_rank']})" if g["rated"] else " (FCS)")
            vals = [f"{t['name']}|{n}", g["wk"], g["opp"], site, g["result"], g.get("expected"), g.get("delucked"), g.get("to"), g.get("resid"),
                    1 if g["rated"] else 0, g.get("opp_rank"), label]
            for j, v in enumerate(vals, 1):
                wg.cell(row=r, column=j, value=v)
            r += 1

    wc = _hidden(wb, "_TapeCurve")
    wc["A1"], wc["B1"] = "margin", "p_home"
    for i, (m, p) in enumerate(d["curve"], 2):
        wc.cell(row=i, column=1, value=m); wc.cell(row=i, column=2, value=p)

    wo = _hidden(wb, "_TapeCommon")
    wo["A1"], wo["B1"] = "pair", "text"
    r = 2
    for a in teams:
        ra = {g["opp_key"]: g for g in a["games"] if g["rated"]}
        for b in teams:
            if a is b:
                continue
            common = [k for k in ra if any(g["opp_key"] == k and g["rated"] for g in b["games"])]
            if not common:
                continue
            parts = []
            for k in common:
                ga = ra[k]; gb = next(g for g in b["games"] if g["opp_key"] == k and g["rated"])
                w = lambda g: {"home": "home", "away": "away", "neutral": "neutral"}[g["site"]]
                parts.append(f"{ga['opp']}: {a['name']} {ga['result']} {w(ga)} (grade {ga['resid']:+.1f}) / {b['name']} {gb['result']} {w(gb)} (grade {gb['resid']:+.1f})")
            wo.cell(row=r, column=1, value=f"{a['name']}|{b['name']}"); wo.cell(row=r, column=2, value="; ".join(parts)); r += 1
    return col


# ---------------- the Matchup sheet ----------------

def write_matchup_sheet(wb, d: dict, col: dict):
    names = sorted(t["name"] for t in d["teams"].values())
    prior = {}
    if "Matchup" in wb.sheetnames:
        old = wb["Matchup"]
        prior = dict(a=old["B1"].value, b=old["B2"].value, site=old["B3"].value)
        del wb["Matchup"]
    idx = 2 if len(wb.sheetnames) > 2 else len(wb.sheetnames)
    ws = wb.create_sheet("Matchup", idx)
    hfa, to_pts, lam, sd, cap = d["rule"]["hfa"], d["rule"]["to_pts"], d["rule"]["lam"], d["rule"]["sd"], d["rule"]["cap"]

    # ---- inputs + hidden helpers (AA list, AB helpers) ----
    ws["A1"], ws["A2"], ws["A3"] = "Team A:", "Team B:", "Site:"
    for c in ("A1", "A2", "A3"):
        ws[c].font = ARIAL_B
    ws["B1"] = prior.get("a") if prior.get("a") in names else "Florida State"
    ws["B2"] = prior.get("b") if prior.get("b") in names else "Miami"
    ws["B3"] = prior.get("site") if prior.get("site") in SITES else "B at home"
    for c in ("B1", "B2", "B3"):
        ws[c].font = Font(name="Arial", bold=True, size=12, color="0000FF"); ws[c].fill = PICK_FILL
    for i, t in enumerate(names, 1):
        ws.cell(row=i, column=27, value=t)                       # AA
    for i, s in enumerate(SITES, 1):
        ws.cell(row=i, column=29, value=s)                       # AC
    dv = DataValidation(type="list", formula1=f"=$AA$1:$AA${len(names)}", allow_blank=False, showDropDown=False)
    ws.add_data_validation(dv); dv.add(ws["B1"]); dv.add(ws["B2"])
    dvs = DataValidation(type="list", formula1="=$AC$1:$AC$3", allow_blank=False, showDropDown=False)
    ws.add_data_validation(dvs); dvs.add(ws["B3"])
    ws["AB1"] = "=MATCH($B$1,_Tape!$B:$B,0)"
    ws["AB2"] = "=MATCH($B$2,_Tape!$B:$B,0)"
    ws["AB3"] = '=IF($B$3="A at home",1,IF($B$3="B at home",-1,0))'
    for c in ("AA", "AB", "AC"):
        ws.column_dimensions[c].hidden = True
    ws["D1"] = f"Tale of the Tape - the machine's read on any two FBS teams. No market numbers. Through Week {d['through_week']}, updated {d['as_of'][:10]}."
    ws["D1"].font = MUTED
    ws["D2"] = "Blue cells are the only inputs. Everything else is a formula off the hidden _Tape sheets (rebuilt by every refresh)."
    ws["D2"].font = MUTED

    A = lambda f: f"INDEX(_Tape!${col[f]}:${col[f]},$AB$1)"       # team A field
    B = lambda f: f"INDEX(_Tape!${col[f]}:${col[f]},$AB$2)"       # team B field
    M = "$B$8"            # margin cell (A minus B)

    # ---- verdict ----
    ws["A5"] = '=$B$1&"  vs  "&$B$2&IF($B$3="Neutral","  (neutral field)",IF($B$3="A at home","  at "&$B$1,"  at "&$B$2))'
    ws["A5"].font = TITLE
    ws["A7"], ws["B7"], ws["C7"] = "", "Team A", "Team B"
    ws["B7"] = "=$B$1"; ws["C7"] = "=$B$2"
    for c in ("B7", "C7"):
        ws[c].font = WHITE_B; ws[c].fill = HEAD_FILL; ws[c].alignment = Alignment(horizontal="center")
    ws["A8"] = "Machine margin (A minus B)"; ws["B8"] = f"={A('now')}-{B('now')}+{hfa}*$AB$3"
    ws["A9"] = "The call"
    ws["B9"] = f'=IF({M}>=0,$B$1&" by "&TEXT({M},"0.0"),$B$2&" by "&TEXT(-{M},"0.0"))'
    ws["B9"].font = BIG
    ws["A10"] = "Win probability"
    ws["B10"] = f"=IF($AB$3=-1,1-VLOOKUP(MAX(-60,MIN(60,-{M})),_TapeCurve!$A:$B,2,TRUE),VLOOKUP(MAX(-60,MIN(60,{M})),_TapeCurve!$A:$B,2,TRUE))"
    ws["C10"] = "=1-B10"
    ws["A11"] = "Score shape (pace guess, not the on-air total)"
    ws["D11"] = f"=({A('off_ppg')}+{B('def_ppg')}+{B('off_ppg')}+{A('def_ppg')})/2"   # expected total, hidden-ish helper
    ws["B11"] = f"=MAX(0,ROUND(($D$11+{M})/2,0))"; ws["C11"] = f"=MAX(0,ROUND(($D$11-{M})/2,0))"
    ws["D11"].font = MUTED; ws["D11"].number_format = '"total "0.0'
    ws["A12"] = "The read"
    gap = f"ABS({A('now')}-{B('now')})"
    ws["B12"] = (f'=IF(ABS({M})<1,"A coin flip: the ratings are "&TEXT({gap},"0.0")&" apart"&IF($AB$3=0," and there is no home field to break it."," and home field all but cancels the gap."),'
                 f'IF(AND({gap}<{hfa},$AB$3<>0),IF({M}>=0,$B$1,$B$2)&" by "&TEXT(ABS({M}),"0.0")&": the two are "&TEXT({gap},"0.0")&" apart on the ratings, so home field ({hfa}) is most of the call.",'
                 f'IF({M}>=0,$B$1,$B$2)&" by "&TEXT(ABS({M}),"0.0")&": a "&TEXT({gap},"0.0")&"-point gap in the ratings"&IF($AB$3=0,"",IF(SIGN({M})*$AB$3>0," plus {hfa} for the home field"," less {hfa} for playing on the road"))&"."))')
    ws.merge_cells("B12:F12"); ws["B12"].alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[12].height = 32
    ws["B8"].number_format = "+0.0;-0.0"; ws["B10"].number_format = "0%"; ws["C10"].number_format = "0%"

    # ---- 1 the number ----
    r = 14
    ws.cell(row=r, column=1, value="1  THE NUMBER").font = WHITE_B
    for c in range(1, 5):
        ws.cell(row=r, column=c).fill = HEAD_FILL
    rows = [("July (ESPN preseason FPI)", "july", "0.0"), ("Now (the machine)", "now", "0.0"), ("Moved since July", "delta", "+0.0;-0.0"),
            ("Rank in the machine", "rank", "0"), ("QB-out adjustment", "qb_out", "0.0"), ("Rated games", "rated_gp", "0"), ("Record", "record", "@")]
    for i, (lab, f, fmt) in enumerate(rows, r + 1):
        ws.cell(row=i, column=1, value=lab).font = ARIAL
        ca = ws.cell(row=i, column=2, value=f"={A(f)}"); cb = ws.cell(row=i, column=3, value=f"={B(f)}")
        for c in (ca, cb):
            c.font = ARIAL; c.number_format = fmt; c.alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add(f"B{r+1}:C{r+2}", DataBarRule(start_type="num", start_value=-10, end_type="num", end_value=35, color="638EC6"))
    jg = f"({A('july')}-{B('july')})"
    ws.cell(row=r + 8, column=1, value="Why").font = ARIAL_B
    ws.cell(row=r + 8, column=2, value=(
        f'="In July the machine had "&IF({jg}>=0,$B$1,$B$2)&" ahead by "&TEXT(ABS({jg}),"0.0")&"; since then "&$B$1&" "&TEXT({A("delta")},"+0.0;-0.0")&", "&$B$2&" "&TEXT({B("delta")},"+0.0;-0.0")'
        f'&" (the July number is worth {lam:g} games of evidence, so after "&{A("rated_gp")}&" and "&{B("rated_gp")}&" rated games it still carries a share)."'
        f'&IF({A("qb_out")}+{B("qb_out")}>0," "&IF({A("qb_out")}>0,$B$1,$B$2)&" is carrying a QB-out adjustment of "&TEXT({A("qb_out")}+{B("qb_out")},"0.0")&" points.","")'))
    ws.merge_cells(start_row=r + 8, start_column=2, end_row=r + 8, end_column=6)
    ws.cell(row=r + 8, column=2).alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[r + 8].height = 44
    ws.cell(row=r + 9, column=1, value="The math").font = ARIAL_B
    ws.cell(row=r + 9, column=2, value=f'={A("abbr")}&" "&TEXT({A("now")},"0.0")&" - "&{B("abbr")}&" "&TEXT({B("now")},"0.0")&IF($AB$3=1," + {hfa} home",IF($AB$3=-1," - {hfa} road"," (neutral)"))&" = "&{A("abbr")}&" "&TEXT({M},"+0.0;-0.0")').font = ARIAL_B

    # ---- 2 the plays (shadow) ----
    r = 25
    ws.cell(row=r, column=1, value="2  THE PLAYS  (SHADOW - the efficiency rating, never the on-air number)").font = WHITE_B
    for c in range(1, 5):
        ws.cell(row=r, column=c).fill = SHADOW_FILL
    rows = [("Offense (points above an average offense)", "eff_off"), ("Defense (points taken away)", "eff_def"), ("Special teams", "eff_st"),
            ("Shadow rating", "eff_total"), ("Shadow rank", "eff_rank")]
    for i, (lab, f) in enumerate(rows, r + 1):
        ws.cell(row=i, column=1, value=lab).font = ARIAL
        for c, F in ((2, A), (3, B)):
            cell = ws.cell(row=i, column=c, value=f"={F(f)}"); cell.font = ARIAL; cell.alignment = Alignment(horizontal="center")
            cell.number_format = "0" if f == "eff_rank" else "+0.0;-0.0"
    aoff = f"({A('eff_off')}-{B('eff_def')})"; boff = f"({B('eff_off')}-{A('eff_def')})"
    em = f"({aoff}-{boff}+({A('eff_st')}-{B('eff_st')})+{hfa}*$AB$3)"
    ws.cell(row=r + 6, column=1, value="A offense vs B defense (net, a game)").font = ARIAL
    ws.cell(row=r + 6, column=2, value=f"={aoff}").number_format = "+0.0;-0.0"
    ws.cell(row=r + 7, column=1, value="B offense vs A defense (net, a game)").font = ARIAL
    ws.cell(row=r + 7, column=2, value=f"={boff}").number_format = "+0.0;-0.0"
    ws.cell(row=r + 8, column=1, value="The plays' margin (A minus B)").font = ARIAL_B
    ws.cell(row=r + 8, column=2, value=f"={em}").number_format = "+0.0;-0.0"
    ws.cell(row=r + 9, column=1, value="50/50 blend with the machine").font = ARIAL
    ws.cell(row=r + 9, column=2, value=f"=({M}+{em})/2").number_format = "+0.0;-0.0"
    ws.cell(row=r + 10, column=1, value="Why").font = ARIAL_B
    dis = f"({em}-{M})"
    AP = "'"
    who_e = f'IF({em}>=0,$B$1,$B$2)'
    who_m = f'IF({M}>=0,$B$1,$B$2)'
    unit = (f'IF(ABS({aoff})>=ABS({boff}),IF({aoff}>=0,$B$1&"{AP}s offense against "&$B$2&"{AP}s defense",$B$2&"{AP}s defense against "&$B$1&"{AP}s offense"),'
            f'IF({boff}>=0,$B$2&"{AP}s offense against "&$B$1&"{AP}s defense",$B$1&"{AP}s defense against "&$B$2&"{AP}s offense"))')
    why = (f'="Play by play "&IF(ABS({dis})<1.5,"agrees with the machine within a point: "&{who_e}&" by "&TEXT(ABS({em}),"0.0")&".",'
           f'IF(SIGN({em})=SIGN({M}),"sides with "&{who_e}&IF(ABS({em})>ABS({M})," harder"," less")&" than the scoreboard does: "&TEXT(ABS({em}),"0.0")&" against the machine{AP}s "&TEXT(ABS({M}),"0.0")&".",'
           f'"leans the other way: "&{who_e}&" by "&TEXT(ABS({em}),"0.0")&" against the machine{AP}s "&{who_m}&" by "&TEXT(ABS({M}),"0.0")&"."))'
           f'&" The bigger unit edge is "&{unit}&" ("&TEXT(MAX(ABS({aoff}),ABS({boff})),"0.0")&" a game)."'
           f'&IF(AND(ABS({dis})>=3,{A("rank")}<=40,{B("rank")}<=40,ABS({M})<=14)," An A&M-type gap: on close games between top-40 teams the final has landed about halfway from the machine toward the plays{AP} number (2022-26).","")')
    ws.cell(row=r + 10, column=2, value=why)
    ws.merge_cells(start_row=r + 10, start_column=2, end_row=r + 10, end_column=6)
    ws.cell(row=r + 10, column=2).alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[r + 10].height = 58

    # ---- 3 the tape ----
    r = 38
    ws.cell(row=r, column=1, value="3  THE TAPE  (season rates per game with FBS ranks; the better-ranked unit takes the edge)").font = WHITE_B
    for c in range(1, 5):
        ws.cell(row=r, column=c).fill = HEAD_FILL
    edge_cells = []
    rr = r + 1
    for O, Dd, oT, dT in ((A, B, "$B$1", "$B$2"), (B, A, "$B$2", "$B$1")):
        ws.cell(row=rr, column=1, value=f'={oT}&" with the ball"').font = ARIAL_B
        ws.cell(row=rr, column=2, value=f'={oT}&" offense"').font = ARIAL_B
        ws.cell(row=rr, column=3, value=f'={dT}&" defense"').font = ARIAL_B
        ws.cell(row=rr, column=4, value="Edge").font = ARIAL_B
        for c in range(1, 5):
            ws.cell(row=rr, column=c).fill = SOFT_FILL
        rr += 1
        for k, lab, _hi in d["stat_keys"]:
            ws.cell(row=rr, column=1, value=lab).font = ARIAL
            ov, orv, dvv, drv = O(f"off_{k}"), O(f"off_{k}_rk"), Dd(f"def_{k}"), Dd(f"def_{k}_rk")
            ws.cell(row=rr, column=2, value=f'=TEXT({ov},"0.0")&"  (#"&{orv}&")"').alignment = Alignment(horizontal="center")
            ws.cell(row=rr, column=3, value=f'=TEXT({dvv},"0.0")&"  (#"&{drv}&")"').alignment = Alignment(horizontal="center")
            e = ws.cell(row=rr, column=4, value=f'=IF({orv}<{drv},{oT},IF({drv}<{orv},{dT},"even"))')
            e.alignment = Alignment(horizontal="center"); e.font = ARIAL
            for c in (2, 3):
                ws.cell(row=rr, column=c).font = ARIAL
            edge_cells.append(f"D{rr}")
            rr += 1
        rr += 1
    rng = f"D{r+2}:D{rr-1}"
    ws.cell(row=rr, column=1, value="Edges").font = ARIAL_B
    ws.cell(row=rr, column=2, value=f'=COUNTIF({rng},$B$1)').alignment = Alignment(horizontal="center")
    ws.cell(row=rr, column=3, value=f'=COUNTIF({rng},$B$2)').alignment = Alignment(horizontal="center")
    ea, eb = f"COUNTIF({rng},$B$1)", f"COUNTIF({rng},$B$2)"
    ws.cell(row=rr + 1, column=1, value="Why").font = ARIAL_B
    ws.cell(row=rr + 1, column=2, value=(
        f'="Row by row the better-ranked unit takes the edge: "&$B$1&" "&{ea}&", "&$B$2&" "&{eb}&". "'
        f'&IF(OR({ea}={eb},SIGN({ea}-{eb})=SIGN({M})),"The tape and the number point the same way.","The tape leans "&IF({ea}>{eb},$B$1,$B$2)&" while the number says "&IF({M}>=0,$B$1,$B$2)&"; the number weighs who the stats came against, the tape does not.")'))
    ws.merge_cells(start_row=rr + 1, start_column=2, end_row=rr + 1, end_column=6)
    ws.cell(row=rr + 1, column=2).alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[rr + 1].height = 32
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["$B$1"], fill=PatternFill("solid", start_color="D9EAD3")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["$B$2"], fill=PatternFill("solid", start_color="FCE5CD")))

    # ---- 4 the resume ----
    r = rr + 4
    ws.cell(row=r, column=1, value="4  THE RESUME  (every result graded against today's ratings: expected, margin with turnover luck out, the difference)").font = WHITE_B
    for c in range(1, 9):
        ws.cell(row=r, column=c).fill = HEAD_FILL
    hdr = ["Wk", "Opponent (machine rank)", "Result", "Expected", "De-lucked", "TO margin", "Grade", ""]
    gcols = {"wk": "B", "opp_label": "L", "result": "E", "expected": "F", "delucked": "G", "to": "H", "resid": "I"}
    helper_col = 30   # AD: the matched row per slot
    for T, tcell, start in (("A", "$B$1", r + 1), ("B", "$B$2", r + SLOTS + 4)):
        ws.cell(row=start, column=1, value=f'={tcell}&"  "&{A("record") if T == "A" else B("record")}').font = ARIAL_B
        for j, h in enumerate(hdr[:7], 1):
            c = ws.cell(row=start + 1, column=j, value=h); c.font = ARIAL_B; c.fill = SOFT_FILL
        for n in range(1, SLOTS + 1):
            row = start + 1 + n
            hc = ws.cell(row=row, column=helper_col, value=f'=IFERROR(MATCH({tcell}&"|{n}",_TapeGames!$A:$A,0),"")')
            h = f"${L(helper_col)}{row}"
            for j, (f, gc) in enumerate([("wk", "B"), ("opp_label", "L"), ("result", "E"), ("expected", "F"), ("delucked", "G"), ("to", "H"), ("resid", "I")], 1):
                c = ws.cell(row=row, column=j, value=f'=IF({h}="","",IF(INDEX(_TapeGames!${gc}:${gc},{h})="","",INDEX(_TapeGames!${gc}:${gc},{h})))')
                c.font = ARIAL
                if f in ("expected", "delucked", "resid"):
                    c.number_format = "+0.0;-0.0"; c.alignment = Alignment(horizontal="center")
                elif f in ("to", "wk"):
                    c.number_format = "+0;-0;0" if f == "to" else "0"; c.alignment = Alignment(horizontal="center")
        ws.conditional_formatting.add(f"G{start+2}:G{start+1+SLOTS}", CellIsRule(operator="greaterThanOrEqual", formula=["7"], fill=PatternFill("solid", start_color="D9EAD3")))
        ws.conditional_formatting.add(f"G{start+2}:G{start+1+SLOTS}", CellIsRule(operator="lessThanOrEqual", formula=["-7"], fill=PatternFill("solid", start_color="F4CCCC")))
        F = A if T == "A" else B
        ws.cell(row=start + SLOTS + 2, column=1, value=(
            f'="Best grade: "&{F("best_opp")}&" ("&TEXT({F("best_resid")},"+0.0;-0.0")&"). Worst: "&{F("worst_opp")}&" ("&TEXT({F("worst_resid")},"+0.0;-0.0")&"). '
            f'Turnover luck so far: "&TEXT({F("luck")},"+0.0;-0.0")&" points of margin. Against expectation over "&{F("rated_n")}&" rated games: "&TEXT({F("resid_sum")},"+0.0;-0.0")&"."')).font = MUTED
        ws.merge_cells(start_row=start + SLOTS + 2, start_column=1, end_row=start + SLOTS + 2, end_column=7)
    ws.column_dimensions[L(helper_col)].hidden = True
    r2 = r + 2 * (SLOTS + 3) + 1
    ws.cell(row=r2, column=1, value="Common opponents").font = ARIAL_B
    ws.cell(row=r2, column=2, value='=IFERROR(INDEX(_TapeCommon!$B:$B,MATCH($B$1&"|"&$B$2,_TapeCommon!$A:$A,0)),"None yet this season.")').font = ARIAL
    ws.merge_cells(start_row=r2, start_column=2, end_row=r2, end_column=7)
    ws.cell(row=r2, column=2).alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[r2].height = 44
    ws.cell(row=r2 + 2, column=1, value=(
        f"How to read it: the number is the on-air machine (July FPI prior, ridge over every rated 2026 game on margin - {to_pts:g} x turnover margin, residual cap {cap:g}, "
        f"prior worth {lam:g} games, home field {hfa}); win probability from the machine's margin curve (sd {sd}); it misses the final margin by about 12.5 points a game. "
        "The plays view is the efficiency SHADOW rating. The tape is per-game season rates with FBS ranks and does not adjust for opponent. No market numbers on this sheet.")).font = MUTED
    ws.merge_cells(start_row=r2 + 2, start_column=1, end_row=r2 + 2, end_column=8)
    ws.cell(row=r2 + 2, column=1).alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[r2 + 2].height = 58

    for c, w in zip("ABCDEFGH", (44, 26, 26, 14, 12, 12, 10, 4)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False


def write_tape_sheet(book: Path = BOOK, rebuild_data: bool = True) -> None:
    if rebuild_data or not DATA.exists():
        import tape_data
        DATA.write_text(json.dumps(tape_data.build(), separators=(",", ":")), encoding="utf-8")
    d = json.loads(DATA.read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(book, keep_vba=book.suffix.lower() == ".xlsm")
    col = write_data_sheets(wb, d)
    write_matchup_sheet(wb, d, col)
    wb.save(book)
    print(f"Matchup sheet written: {len(d['teams'])} teams, through week {d['through_week']}")


if __name__ == "__main__":
    from refresh_all import recalc_com, wait_for_unlock
    wait_for_unlock(BOOK)
    write_tape_sheet(BOOK, rebuild_data="--no-data" not in sys.argv)
    if "--no-recalc" not in sys.argv:
        recalc_com(BOOK)
